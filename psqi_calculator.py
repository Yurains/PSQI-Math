# -*- coding: utf-8 -*-
"""
匹茲堡睡眠品質量表 (PSQI) 分數計算機
使用 PyQt5 實作

計分方式（Buysse et al., 1989）：
  C1 主觀睡眠品質   = 第6題
  C2 睡眠潛伏期     = 第2題(分鐘換算) + 第5(1)題 → 0 / 1-2 / 3-4 / 5-6 → 0~3
  C3 睡眠時數       = 第4題：≥7h=0, 6~<7h=1, 5~<6h=2, <5h=3
  C4 睡眠效率       = 第4題 ÷ (第3題 − 第1題) × 100%：≥85=0, 75~84=1, 65~74=2, <65=3
  C5 睡眠困擾       = 第5(2)~(10)題加總：0 / 1-9 / 10-18 / 19-27 → 0~3
  C6 安眠藥物使用   = 第7題
  C7 日間功能障礙   = 第8題 + 第9題 → 0 / 1-2 / 3-4 / 5-6 → 0~3
  總分 0~21，> 5 分表示睡眠品質不佳

第5(10)題「其他」為非必填，未作答時以 0 分計。
"""

import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PyQt5.QtCore import Qt, QDate, QTime
from PyQt5.QtGui import QFont
from PyQt5.QtWidgets import (
    QApplication, QButtonGroup, QFrame, QGridLayout, QGroupBox, QHBoxLayout,
    QDateEdit, QFileDialog, QLabel, QLineEdit, QMainWindow, QMessageBox, QPushButton, QRadioButton,
    QScrollArea, QSpinBox, QTableWidget, QTableWidgetItem, QTimeEdit,
    QVBoxLayout, QWidget, QHeaderView,
)

FREQ_OPTIONS = ["從未發生", "不到一次", "約一兩次", "三次或以上"]

Q5_ITEMS = [
    "(1) 無法在 30 分鐘內入睡",
    "(2) 半夜或凌晨便清醒",
    "(3) 必須起來上廁所",
    "(4) 覺得呼吸不順暢",
    "(5) 大聲打鼾或咳嗽",
    "(6) 會覺得冷",
    "(7) 覺得躁熱",
    "(8) 作惡夢",
    "(9) 身上有疼痛",
    "(10) 其他",
]

Q6_OPTIONS = ["很好", "還不錯", "差了點", "很差"]
Q7_OPTIONS = ["未發生", "不到一次", "一兩次", "三次或三次以上"]
Q8_OPTIONS = ["未發生", "不到一次", "一兩次", "三次或三次以上"]
Q9_OPTIONS = ["沒有", "有一點", "的確有", "很嚴重"]

COMPONENT_NAMES = [
    "C1 主觀睡眠品質",
    "C2 睡眠潛伏期",
    "C3 睡眠時數",
    "C4 睡眠效率",
    "C5 睡眠困擾",
    "C6 安眠藥物使用",
    "C7 日間功能障礙",
]

EXPORT_HEADERS = (
    ["Date", "ID", "Time", 1, "2a", "2b", 3, 4]
    + [f"5{letter}" for letter in "abcdefghij"]
    + [6, 7, 8, 9]
    + [
        "Sleep quality", "Sleep latency", "Sleep duration", "Sleep efficiency",
        "Sleep disturbance", "Use of sleep medication", "Daytime dysfunction",
        "PSQI Score",
    ]
)


def save_psqi_record(path, record):
    """Save one completed questionnaire as a row, appending to a compatible workbook."""
    try:
        workbook = load_workbook(path)
    except FileNotFoundError:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "PSQI 資料"
        sheet.append(EXPORT_HEADERS)
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_HEADERS))}1"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="315A81")
            cell.alignment = Alignment(horizontal="center")
        for index in range(1, len(EXPORT_HEADERS) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 20 if index > 3 else 16
    else:
        sheet = workbook.active
        existing = [sheet.cell(1, col).value for col in range(1, len(EXPORT_HEADERS) + 1)]
        if existing != EXPORT_HEADERS or sheet.max_column != len(EXPORT_HEADERS):
            raise ValueError("所選 Excel 檔案的欄位與 PSQI 匯出格式不同，請選擇新檔名。")

    sheet.append(record)
    row = sheet.max_row
    sheet.cell(row, 1).number_format = "yyyy/mm/dd"
    for col in (4, 7):
        sheet.cell(row, col).number_format = "hh:mm:ss AM/PM"
    workbook.save(path)


# ---------------------------------------------------------------- 計分邏輯 --

def _sum_to_score(total):
    """0→0, 1-2→1, 3-4→2, 5-6→3（C2、C7 共用）"""
    if total == 0:
        return 0
    if total <= 2:
        return 1
    if total <= 4:
        return 2
    return 3


def compute_psqi(bed_min, latency_min, wake_min, sleep_min, q5, q6, q7, q8, q9):
    """
    bed_min / wake_min：上床、起床時間（自午夜起算的分鐘數）
    latency_min：入睡所需分鐘數
    sleep_min：實際睡眠分鐘數
    q5：長度 10 的串列（0~3）；q6~q9：0~3
    回傳 (components, total, detail)
    """
    # C1
    c1 = q6

    # C2
    if latency_min <= 15:
        lat = 0
    elif latency_min <= 30:
        lat = 1
    elif latency_min <= 60:
        lat = 2
    else:
        lat = 3
    c2 = _sum_to_score(lat + q5[0])

    # C3
    hours = sleep_min / 60
    if hours >= 7:
        c3 = 0
    elif hours >= 6:
        c3 = 1
    elif hours >= 5:
        c3 = 2
    else:
        c3 = 3

    # C4（跨午夜處理）
    in_bed = (wake_min - bed_min) % (24 * 60)
    if in_bed == 0:
        in_bed = 24 * 60
    efficiency = sleep_min / in_bed * 100
    if efficiency >= 85:
        c4 = 0
    elif efficiency >= 75:
        c4 = 1
    elif efficiency >= 65:
        c4 = 2
    else:
        c4 = 3

    # C5
    dist = sum(q5[1:])
    if dist == 0:
        c5 = 0
    elif dist <= 9:
        c5 = 1
    elif dist <= 18:
        c5 = 2
    else:
        c5 = 3

    # C6
    c6 = q7

    # C7
    c7 = _sum_to_score(q8 + q9)

    components = [c1, c2, c3, c4, c5, c6, c7]
    detail = {
        "latency_sub": lat,
        "in_bed_min": in_bed,
        "efficiency": efficiency,
        "disturb_sum": dist,
        "daytime_sum": q8 + q9,
    }
    return components, sum(components), detail


# -------------------------------------------------------------------- 介面 --

class PSQIWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("匹茲堡睡眠品質量表 (PSQI) 分數計算機")
        self.resize(900, 820)

        content = QWidget()
        layout = QVBoxLayout(content)
        layout.setSpacing(12)

        title = QLabel("匹茲堡睡眠品質量表")
        title.setFont(QFont("Microsoft JhengHei", 16, QFont.Bold))
        title.setAlignment(Qt.AlignCenter)
        layout.addWidget(title)
        hint = QLabel("請就過去一個月來的日常（大多數）睡眠習慣作答。"
                      "<span style='color:#c0392b'>*</span> 為必填。")
        hint.setAlignment(Qt.AlignCenter)
        layout.addWidget(hint)

        layout.addWidget(self._build_record_info())
        layout.addWidget(self._build_q1_4())
        layout.addWidget(self._build_q5())
        self.q6 = self._build_single("6. 整體而言，你覺得自己的睡眠品質如何？ *", Q6_OPTIONS, layout)
        self.q7 = self._build_single("7. 通常一星期幾個晚上需要使用藥物幫忙睡眠？ *", Q7_OPTIONS, layout)
        self.q8 = self._build_single("8. 是否曾在用餐、開車或社交場合瞌睡而無法保持清醒，每星期約幾次？ *",
                                     Q8_OPTIONS, layout)
        self.q9 = self._build_single("9. 你會感到無心完成該做的事 *", Q9_OPTIONS, layout)
        layout.addStretch()

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setWidget(content)

        # 右側結果區
        result_panel = self._build_result_panel()

        central = QWidget()
        main = QHBoxLayout(central)
        main.addWidget(scroll, 3)
        main.addWidget(result_panel, 2)
        self.setCentralWidget(central)

    # ---- 題目區塊 ----

    def _build_record_info(self):
        box = QGroupBox("匯出資料")
        row = QHBoxLayout(box)
        self.record_date = QDateEdit(QDate.currentDate())
        self.record_date.setCalendarPopup(True)
        self.record_date.setDisplayFormat("yyyy-MM-dd")
        row.addWidget(QLabel("Date"))
        row.addWidget(self.record_date)
        row.addWidget(QLabel("ID、Time 於匯出後在 Excel 填寫"))
        return box

    def _build_q1_4(self):
        box = QGroupBox("基本睡眠資料")
        grid = QGridLayout(box)

        self.bed_time = QTimeEdit(QTime(23, 0))
        self.bed_time.setDisplayFormat("HH:mm")
        self.latency = QSpinBox()
        self.latency.setRange(0, 600)
        self.latency.setSuffix(" 分鐘")
        self.latency.setValue(15)
        self.wake_time = QTimeEdit(QTime(7, 0))
        self.wake_time.setDisplayFormat("HH:mm")
        self.sleep_h = QSpinBox()
        self.sleep_h.setRange(0, 24)
        self.sleep_h.setSuffix(" 時")
        self.sleep_h.setValue(7)
        self.sleep_m = QSpinBox()
        self.sleep_m.setRange(0, 59)
        self.sleep_m.setSuffix(" 分")

        sleep_row = QHBoxLayout()
        sleep_row.addWidget(self.sleep_h)
        sleep_row.addWidget(self.sleep_m)
        sleep_row.addStretch()

        grid.addWidget(QLabel("1. 通常何時上床？ *"), 0, 0)
        grid.addWidget(self.bed_time, 0, 1)
        grid.addWidget(QLabel("2. 通常多久才能入睡？ *"), 1, 0)
        grid.addWidget(self.latency, 1, 1)
        grid.addWidget(QLabel("3. 早上通常何時起床？ *"), 2, 0)
        grid.addWidget(self.wake_time, 2, 1)
        grid.addWidget(QLabel("4. 實際每晚可以入睡幾小時？ *"), 3, 0)
        grid.addLayout(sleep_row, 3, 1)
        grid.setColumnStretch(2, 1)
        return box

    def _build_q5(self):
        box = QGroupBox("5. 睡眠出現下列困擾情形，每星期約有幾次？")
        grid = QGridLayout(box)
        for col, text in enumerate(FREQ_OPTIONS, start=1):
            lbl = QLabel(text)
            lbl.setAlignment(Qt.AlignCenter)
            grid.addWidget(lbl, 0, col)
        self.q5 = []
        for row, item in enumerate(Q5_ITEMS, start=1):
            # 5(10)「其他」為選填，其餘必填
            label = item + "（選填）" if row == len(Q5_ITEMS) else item + " *"
            grid.addWidget(QLabel(label), row, 0)
            group = QButtonGroup(box)
            for col in range(4):
                rb = QRadioButton()
                group.addButton(rb, col)
                cell = QHBoxLayout()
                cell.setAlignment(Qt.AlignCenter)
                cell.addWidget(rb)
                grid.addLayout(cell, row, col + 1)
            self.q5.append(group)
        grid.setColumnStretch(0, 1)

        self.q5_other = QLineEdit()
        self.q5_other.setPlaceholderText("（10）其他，請說明（選填）")
        grid.addWidget(self.q5_other, len(Q5_ITEMS) + 1, 0, 1, 5)
        return box

    def _build_single(self, title, options, layout):
        box = QGroupBox(title)
        row = QHBoxLayout(box)
        group = QButtonGroup(box)
        for i, text in enumerate(options):
            rb = QRadioButton(text)
            group.addButton(rb, i)
            row.addWidget(rb)
        row.addStretch()
        layout.addWidget(box)
        return group

    # ---- 結果區 ----

    def _build_result_panel(self):
        panel = QFrame()
        v = QVBoxLayout(panel)

        btn_row = QHBoxLayout()
        calc_btn = QPushButton("計算分數")
        calc_btn.setMinimumHeight(40)
        calc_btn.setStyleSheet("font-size:15px; font-weight:bold;")
        calc_btn.clicked.connect(self.calculate)
        reset_btn = QPushButton("清除重填")
        reset_btn.setMinimumHeight(40)
        reset_btn.clicked.connect(self.reset)
        export_btn = QPushButton("匯出 Excel")
        export_btn.setMinimumHeight(40)
        export_btn.clicked.connect(self.export_excel)
        btn_row.addWidget(calc_btn, 2)
        btn_row.addWidget(reset_btn, 1)
        btn_row.addWidget(export_btn, 1)
        v.addLayout(btn_row)

        self.table = QTableWidget(len(COMPONENT_NAMES), 3)
        self.table.setHorizontalHeaderLabels(["構面", "分數", "說明"])
        self.table.verticalHeader().setVisible(False)
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.ResizeToContents)
        header.setSectionResizeMode(1, QHeaderView.ResizeToContents)
        header.setSectionResizeMode(2, QHeaderView.Stretch)
        for r, name in enumerate(COMPONENT_NAMES):
            self.table.setItem(r, 0, QTableWidgetItem(name))
            self.table.setItem(r, 1, QTableWidgetItem("-"))
            self.table.setItem(r, 2, QTableWidgetItem(""))
        v.addWidget(self.table)

        self.total_label = QLabel("總分：-")
        self.total_label.setFont(QFont("Microsoft JhengHei", 20, QFont.Bold))
        self.total_label.setAlignment(Qt.AlignCenter)
        v.addWidget(self.total_label)

        self.verdict_label = QLabel("")
        self.verdict_label.setAlignment(Qt.AlignCenter)
        self.verdict_label.setWordWrap(True)
        self.verdict_label.setFont(QFont("Microsoft JhengHei", 12))
        v.addWidget(self.verdict_label)

        self.warn_label = QLabel("")
        self.warn_label.setWordWrap(True)
        self.warn_label.setStyleSheet("color:#d35400;")
        v.addWidget(self.warn_label)

        note = QLabel("總分範圍 0～21 分；總分 > 5 分表示睡眠品質不佳。")
        note.setWordWrap(True)
        note.setStyleSheet("color:gray;")
        v.addWidget(note)
        v.addStretch()
        return panel

    # ---- 動作 ----

    def calculate(self):
        # 5(10)「其他」為選填，不檢查
        missing = [Q5_ITEMS[i] for i, g in enumerate(self.q5[:-1]) if g.checkedId() < 0]
        missing = ["第5題 " + m for m in missing]
        for name, g in (("第6題", self.q6), ("第7題", self.q7), ("第8題", self.q8), ("第9題", self.q9)):
            if g.checkedId() < 0:
                missing.append(name)
        sleep_min = self.sleep_h.value() * 60 + self.sleep_m.value()
        if sleep_min == 0:
            missing.append("第4題（實際睡眠時數不可為 0）")
        if missing:
            QMessageBox.warning(self, "尚未完成", "請完成以下必填項目：\n\n" + "\n".join(missing))
            return None

        bed = self.bed_time.time()
        wake = self.wake_time.time()
        components, total, d = compute_psqi(
            bed_min=bed.hour() * 60 + bed.minute(),
            latency_min=self.latency.value(),
            wake_min=wake.hour() * 60 + wake.minute(),
            sleep_min=sleep_min,
            q5=[max(g.checkedId(), 0) for g in self.q5],  # 未作答（選填項）以 0 計
            q6=self.q6.checkedId(),
            q7=self.q7.checkedId(),
            q8=self.q8.checkedId(),
            q9=self.q9.checkedId(),
        )

        in_bed_h = d["in_bed_min"] / 60
        notes = [
            f"第6題：{Q6_OPTIONS[self.q6.checkedId()]}",
            f"入睡 {self.latency.value()} 分（{d['latency_sub']}）+ 第5(1)題（{self.q5[0].checkedId()}）"
            f" = {d['latency_sub'] + self.q5[0].checkedId()}",
            f"實際睡眠 {sleep_min / 60:.2f} 小時",
            f"{sleep_min / 60:.2f} ÷ 在床 {in_bed_h:.2f} 小時 = {d['efficiency']:.1f}%",
            f"第5(2)~(10)題加總 = {d['disturb_sum']}",
            f"第7題：{Q7_OPTIONS[self.q7.checkedId()]}",
            f"第8題 + 第9題 = {d['daytime_sum']}",
        ]
        for r, (score, note) in enumerate(zip(components, notes)):
            self.table.item(r, 1).setText(str(score))
            self.table.item(r, 1).setTextAlignment(Qt.AlignCenter)
            self.table.item(r, 2).setText(note)

        self.total_label.setText(f"總分：{total} / 21")
        if total > 5:
            self.verdict_label.setText("睡眠品質不佳（> 5 分）")
            self.verdict_label.setStyleSheet("color:#c0392b; font-weight:bold;")
        else:
            self.verdict_label.setText("睡眠品質良好（≤ 5 分）")
            self.verdict_label.setStyleSheet("color:#27ae60; font-weight:bold;")

        warns = []
        if d["efficiency"] > 100:
            warns.append("注意：實際睡眠時數大於在床時間（效率 > 100%），請確認第1、3、4題是否填寫正確。")
        self.warn_label.setText("\n".join(warns))
        return components, total

    def export_excel(self):
        result = self.calculate()
        if result is None:
            return
        components, total = result
        latency_min = self.latency.value()
        latency_score = 0 if latency_min <= 15 else 1 if latency_min <= 30 else 2 if latency_min <= 60 else 3

        path, _ = QFileDialog.getSaveFileName(
            self, "匯出 PSQI 資料", "PSQI_資料.xlsx", "Excel 活頁簿 (*.xlsx)"
        )
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        record = [
            self.record_date.date().toPyDate(), None,
            None,
            self.bed_time.time().toPyTime(), latency_min, latency_score,
            self.wake_time.time().toPyTime(),
            self.sleep_h.value() + self.sleep_m.value() / 60,
            *[max(g.checkedId(), 0) for g in self.q5],
            self.q6.checkedId(), self.q7.checkedId(), self.q8.checkedId(), self.q9.checkedId(),
            *components, total,
        ]
        try:
            save_psqi_record(path, record)
        except (OSError, ValueError) as exc:
            QMessageBox.critical(self, "匯出失敗", str(exc))
            return
        QMessageBox.information(self, "匯出完成", f"已儲存至：\n{path}")

    def reset(self):
        self.record_date.setDate(QDate.currentDate())
        self.bed_time.setTime(QTime(23, 0))
        self.wake_time.setTime(QTime(7, 0))
        self.latency.setValue(15)
        self.sleep_h.setValue(7)
        self.sleep_m.setValue(0)
        for g in self.q5 + [self.q6, self.q7, self.q8, self.q9]:
            g.setExclusive(False)
            for b in g.buttons():
                b.setChecked(False)
            g.setExclusive(True)
        self.q5_other.clear()
        for r in range(len(COMPONENT_NAMES)):
            self.table.item(r, 1).setText("-")
            self.table.item(r, 2).setText("")
        self.total_label.setText("總分：-")
        self.verdict_label.setText("")
        self.warn_label.setText("")


def main():
    app = QApplication(sys.argv)
    app.setFont(QFont("Microsoft JhengHei", 10))
    win = PSQIWindow()
    win.show()
    sys.exit(app.exec_())


if __name__ == "__main__":
    main()
